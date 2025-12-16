from rest_framework.viewsets import GenericViewSet
from rest_framework import mixins, viewsets
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework import serializers

import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from django.http import HttpResponse
from django.contrib.auth.models import User
from django.contrib.auth import authenticate, login, logout
from django.db.models import Count, Q

from dotaesports.models import Tournament, Team, Player, TournamentTeamParticipation, Match, PlayerTeamHistory
from dotaesports.serializers import TournamentSerializer, TeamSerializer, PlayerSerializer, TournamentTeamParticipationSerializer, MatchSerializer, PlayerTeamHistorySerializer, UserSerializer

class TournamentViewset(
    mixins.CreateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.UpdateModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
):
    queryset = Tournament.objects.all()
    serializer_class = TournamentSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        
        if not self.request.user.is_staff and self.request.user.id:
            user_player = Player.objects.get(user=self.request.user)
            user_team = user_player.team
            
            team_tournament_ids = TournamentTeamParticipation.objects.filter(team=user_team).values_list('tournament_id',flat=True)
            qs = qs.filter(id__in=team_tournament_ids)

        return qs
    
    @action(detail=False, methods=['GET'], url_path='export-excel')
    def export_tournaments_to_excel(self, request, *args, **kwargs):
        tournaments = Tournament.objects.all()

        wb = Workbook()
        ws = wb.active
        ws.title = "Турниры"
        
        headers = ['Название', 'Дата начала', 'Дата окончания', 'Статус', 'Призовой фонд', 'Логотип']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        for row, tournament in enumerate(tournaments, 2):
            ws.cell(row=row, column=1, value=tournament.name)
            ws.cell(row=row, column=2, value=tournament.start_date)
            ws.cell(row=row, column=3, value=tournament.end_date)

            status_map = {
                'TBA': 'Неизвестно',
                'Coming': 'Скоро начнётся', 
                'Live': 'Идёт',
                'Ended': 'Закончен'
            }
            status_text = status_map.get(tournament.status, tournament.status)
            ws.cell(row=row, column=4, value=status_text)
            
            ws.cell(row=row, column=5, value=tournament.prize_pool or 'Не определён')
            ws.cell(row=row, column=6, value='Да' if tournament.logo else 'Нет')

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="tournaments.xlsx"'
        
        return response

    @action(detail=True, methods=['GET'], url_path="team_count_stats")
    def team_count(self, request, pk=None):
        tournament = self.get_object()
        
        count = TournamentTeamParticipation.objects.filter(tournament=tournament).aggregate(team_count=Count('team', distinct=True))
        
        return Response({'count': count.get('team_count', 0)})
    
class TeamViewset(
    mixins.CreateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.UpdateModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
):
    queryset = Team.objects.all()
    serializer_class = TeamSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        
        if not self.request.user.is_staff and self.request.user.id:
            user_player = Player.objects.get(user=self.request.user)
            qs = qs.filter(id=user_player.team.id)
        
        return qs

    class PrizeStatsSerializer(serializers.Serializer):
        first_places = serializers.IntegerField()
        second_places = serializers.IntegerField()
        third_places = serializers.IntegerField()
    
    @action(detail=True, methods=['GET'], url_path="prize_stats")
    def prize_stats(self, request, *args, **kwargs):
        team = self.get_object()
        
        stats = TournamentTeamParticipation.objects.filter(team=team).aggregate(
            first=Count('id', filter=Q(place=1)),
            second=Count('id', filter=Q(place=2)),
            third=Count('id', filter=Q(place=3))
        )
        
        return Response({
            'first': stats.get('first', 0),
            'second': stats.get('second', 0),
            'third': stats.get('third', 0)
        })
    


class PlayerViewset(
    mixins.CreateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.UpdateModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
):
    queryset = Player.objects.all()
    serializer_class = PlayerSerializer

    def get_queryset(self):
        qs = super().get_queryset()

        user_team = Team()

        for player in qs:
            if player.user == self.request.user:
                user_team = player.team
                break

        if not self.request.user.is_staff and self.request.user.id:
            qs = qs.filter(team=user_team)

        return qs
    
    @action(detail=True, methods=['GET'], url_path="team_history")
    def team_history(self, request, *args, **kwargs):
        player = self.get_object()
        
        history = PlayerTeamHistory.objects.filter(player=player).order_by('-created_at')
        
        data = []
        for record in history:
            data.append({
                'team': record.team.name if record.team else 'Без команды',
                'date': record.created_at.strftime('%d.%m.%Y %H:%M') if record.created_at else '',
                'manager': record.manager.username if record.manager else 'Неизвестно'
            })
        
        return Response(data)
    
    @action(detail=False, methods=['GET'], url_path='export-excel')
    def export_players_to_excel(self, request, *args, **kwargs):
        players = Player.objects.all()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Игроки"
        
        headers = ['Никнейм', 'Настоящее имя', 'Роль', 'Команда', 'Фото']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
        
        role_map = {
            'CARRY': 'Керри',
            'MIDLANER': 'Мидлейнер',
            'HARDLINER': 'Тройка',
            'SEMISUPPORT': 'Четвёрка',
            'FULLSUPPORT': 'Пятёрка',
        }
        
        for row, player in enumerate(players, 2):
            ws.cell(row=row, column=1, value=player.nickname)
            ws.cell(row=row, column=2, value=player.real_name or '')
            ws.cell(row=row, column=3, value=role_map.get(player.role, player.role))
            ws.cell(row=row, column=4, value=player.team.name if player.team else 'Без команды')
            ws.cell(row=row, column=5, value='Да' if player.photo else 'Нет')
        
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="players.xlsx"'
        
        return response
    


class TournamentTeamParticipationViewset(
    mixins.CreateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.UpdateModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
):
    queryset = TournamentTeamParticipation.objects.all()
    serializer_class = TournamentTeamParticipationSerializer

    def get_queryset(self):
        qs = super().get_queryset()

        if not self.request.user.is_staff and self.request.user.id:
            user_player = Player.objects.get(user=self.request.user)
            qs = qs.filter(team=user_player.team)
        
        return qs

class MatchViewset(
    mixins.CreateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.UpdateModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
):
    queryset = Match.objects.all()
    serializer_class = MatchSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        
        if not self.request.user.is_staff and self.request.user.id:
            user_player = Player.objects.get(user=self.request.user)
            user_team = user_player.team

            qs = qs.filter(radiant=user_team) | qs.filter(dire=user_team)
        
        return qs

class UserViewset(mixins.CreateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.UpdateModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet):
    queryset = User.objects.all()
    serializer_class = UserSerializer

    @action(url_path="my", methods=["GET"], detail=False)
    def get_my(self, *args, **kwargs):
        return Response({
            'username': self.request.user.username,
            'is_authenticated': self.request.user.is_authenticated,
            'is_staff': self.request.user.is_staff
        })
    
    @action(url_path="login", methods=["POST"], detail=False)
    def process_login(self, *args, **kwargs):
        class LoginSerializer(serializers.Serializer):
            username = serializers.CharField()
            password = serializers.CharField()

        serializer = LoginSerializer(data=self.request.data)
        serializer.is_valid(raise_exception=True)

        username = serializer.validated_data['username']
        password = serializer.validated_data['password']
        
        user = authenticate(username=username, password=password)
        if user:
            login(request=self.request, user=user)
        else:
            return Response({
                "status": "Failed!"
            }, status=401)
        return Response({
                "status": "Success!"
        })
    
    @action(url_path="logout", methods=["POST"], detail=False)
    def process_logout(self, *args, **kwargs):
        logout(self.request)

        return Response({
                "status": "Success!"
        })

class PlayerTeamHistoryViewset(
    mixins.CreateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.UpdateModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
):
    queryset = PlayerTeamHistory.objects.all()
    serializer_class = PlayerTeamHistorySerializer
    
    def get_queryset(self):
        qs = super().get_queryset()

        if not self.request.user.is_staff and self.request.user.id:
            try:
                user_player = Player.objects.get(user=self.request.user)
                qs = qs.filter(player__team=user_player.team)
            except Player.DoesNotExist:
                qs = qs.none()
        
        return qs